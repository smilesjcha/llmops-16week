"""Create the editable 8-week proposal Word template (PDF exported separately)."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TARGET = ROOT / "week08/templates/08_project_proposal_word_template.docx"
NAVY = RGBColor(16, 39, 71)
GRAY = RGBColor(83, 97, 116)
LINE = "D4DAE3"


def set_cell_bg(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def para(doc, text="", *, size=10.5, bold=False, color=None, space=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.line_spacing = 1.16
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Apple SD Gothic Neo"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return p


def heading(doc, n, title, prompt):
    para(doc, f"{n}  {title}", size=12.5, bold=True, color=NAVY, space=2)
    para(doc, prompt, size=9.5, color=GRAY, space=5)


def grid(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    if widths:
        for i, width in enumerate(widths):
            table.columns[i].width = Cm(width)
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        c.text = h
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_bg(c, "102747")
        for r in c.paragraphs[0].runs:
            r.bold = True
            r.font.name = "Apple SD Gothic Neo"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(table.rows) % 2 == 0:
                set_cell_bg(cells[i], "F6F7F9")
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.name = "Apple SD Gothic Neo"
                    r.font.size = Pt(9)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


doc = Document()
sec = doc.sections[0]
sec.page_width = Cm(21.0)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.5)
sec.bottom_margin = Cm(1.4)
sec.left_margin = Cm(1.8)
sec.right_margin = Cm(1.8)
sec.header_distance = Cm(0.8)
sec.footer_distance = Cm(0.8)

style = doc.styles["Normal"]
style.font.name = "Apple SD Gothic Neo"
style.font.size = Pt(10.5)
style.font.color.rgb = RGBColor(23, 25, 29)
style.paragraph_format.space_after = Pt(5)

h = sec.header.paragraphs[0]
h.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = h.add_run("LLMOPS  /  2026-2  /  PROJECT PROPOSAL")
r.font.name = "Apple SD Gothic Neo"
r.font.size = Pt(8)
r.font.color.rgb = GRAY

para(doc, "서비스 기획서 · 2Pager", size=21, bold=True, color=NAVY, space=4)
para(doc, "초안 0.1 → 교수 피드백 → 수정본 0.2", size=10, color=GRAY, space=10)
grid(doc, ["프로젝트", "작성 정보", "버전"], [["[한 문장 서비스명]", "[개인 또는 팀명 · 소속 · 연락처]", "[0.1 또는 0.2]"]], [6.2, 8.0, 2.8])

heading(doc, "01", "제안 요약", "첫 문단에 누구의 어떤 문제와 해결 가치가 있는지 적습니다.")
para(doc, "[한 문장 제안: 대상 사용자 + 현재 불편 + 제안하는 서비스의 핵심 가치]", space=8)

heading(doc, "02", "문제와 근거", "관찰한 사실, 팀의 가정, 이후 검증할 가설을 분리합니다.")
grid(doc, ["구분", "기록할 내용"], [
    ["현재 방식", "[사용자가 지금 무엇을 어떻게 해결하는지]"],
    ["확인한 근거", "[원문 URL·조사일·인터뷰/공개 자료 범위]"],
    ["검증할 가설", "[아직 확인하지 않은 사용자 요구·성과 가정]"],
], [4.0, 13.0])

heading(doc, "03", "사용자 흐름과 범위", "대기·오류·답변 보류·수정 경로까지 포함합니다.")
para(doc, "[입력 → 핵심 처리 → 결과 확인 → 사용자 수정/재시도]", space=4)
grid(doc, ["이번 프로젝트의 최소 기능", "제외 범위"], [
    ["[핵심 기능 1·2·3과 완료 조건]", "[이번 학기에 구현하지 않을 기능·자동화]"],
], [8.5, 8.5])

heading(doc, "04", "핵심 요구사항", "요구사항을 구현 여부가 판정되는 문장으로 적습니다.")
grid(doc, ["ID", "요구사항", "완료 조건"], [
    ["R01", "[사용자에게 제공할 기능]", "[입력·출력·실패 조건]"],
    ["R02", "[근거·출처·버전 표시]", "[검증 가능한 결과]"],
    ["R03", "[오류·보류·수정 경로]", "[사용자가 취할 다음 행동]"],
], [1.6, 7.2, 8.2])

doc.add_page_break()
para(doc, "기술·검증·실행 계획", size=18, bold=True, color=NAVY, space=8)
heading(doc, "05", "데이터와 기술 구조", "실제 데이터의 사용 권한, 보관·삭제 범위, 입력·출력을 명시합니다.")
grid(doc, ["결정 항목", "현재 설계와 근거"], [
    ["데이터", "[출처·라이선스/동의·버전·민감정보 여부]"],
    ["처리 흐름", "[입력 → 검색/모델/도구 → 검증 → 결과]"],
    ["운영 보호", "[접근 권한·로그·삭제·사람 검토·복구]"],
], [4.0, 13.0])

heading(doc, "06", "평가 계획", "기준선과 비교 조건을 먼저 고정하고 성공·실패·보류를 분리합니다.")
grid(doc, ["평가 항목", "사전에 정할 기준"], [
    ["기준선", "[현재 방식 또는 단순한 모델/검색 설정]"],
    ["대표 입력", "[답이 있는 사례·없는 사례·표현 차이·권한/날짜 사례]"],
    ["성공·실패", "[품질 지표와 허용 오차·대표 실패 2건의 분석 방법]"],
    ["운영 지표", "[응답 시간·비용·오류율의 측정 조건]"],
], [4.0, 13.0])

heading(doc, "07", "일정과 역할", "7주차 초안에서 확인할 질문, 8주차 수정 이유, 16주차 시연 범위를 적습니다.")
grid(doc, ["시점", "산출물·결정"], [
    ["7주차 · 10.16", "초안 0.1 제출 · 교수 피드백 기록"],
    ["8주차 · 10.23", "피드백 반영 수정본 0.2 온라인 제출·평가"],
    ["9–15주차", "[구현·평가·운영 검증의 담당과 완료 기준]"],
    ["16주차", "실시간 온라인 발표 · [시연 범위와 질문]"],
], [4.0, 13.0])

heading(doc, "08", "검토 요청과 수정 이력", "사업 가치·사용 경험·구현 가능성 관점에서 확인할 질문을 적습니다.")
para(doc, "[교수에게 묻고 싶은 질문 2–3개]", space=3)
para(doc, "[수정본에서는 피드백 항목 → 변경 내용 → 판단 근거를 기록]", size=9.5, color=GRAY, space=9)

para(doc, "제출 안내  아주대AI대학원 팀: LMS · 다른 학교 개인: 공지된 교수 이메일", size=9, bold=True, color=NAVY, space=1)
para(doc, "PPT·Word·Word 기반 PDF·독립형 HTML 가능 · 한글(HWP) 제외 · 정확한 마감 시각은 소속별 공지", size=8.5, color=GRAY)

foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = foot.add_run("2026-2  ·  PROJECT PROPOSAL TEMPLATE")
fr.font.name = "Apple SD Gothic Neo"
fr.font.size = Pt(8)
fr.font.color.rgb = GRAY

TARGET.parent.mkdir(parents=True, exist_ok=True)
doc.save(TARGET)
print(TARGET)
